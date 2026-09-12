"""Loss-conscious intake. Draft records are valid intake, not validated chemistry."""

import csv
import json
import math
from datetime import date, datetime
from pathlib import Path

from .schema import TABLES, VERSION
from .localization import EN, EN_SHEETS, canonical_choice


def clean(value):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return (value.strip() or None) if isinstance(value, str) else value


def read_dataset(path):
    path = Path(path)
    if path.suffix == ".json":
        return json.loads(path.read_text(encoding="utf-8"))
    result = {"schema_version": VERSION}
    if path.is_dir():
        for name in TABLES:
            with (path / f"{name}.csv").open(encoding="utf-8-sig", newline="") as fh:
                reader = csv.reader(fh)
                result[name] = _rows(name, list(reader))
    elif path.suffix.lower() == ".xlsx":
        from openpyxl import load_workbook
        # Formulas are rejected rather than silently importing stale cached values.
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            for name, spec in TABLES.items():
                matching = [s for s in (spec['sheet'], EN_SHEETS[name]) if s in wb.sheetnames]
                if len(matching) > 1:
                    raise ValueError(f"Ambiguous duplicate input sheets for {name}; merge separate workbooks instead")
                if not matching:
                    raise ValueError(f"缺少工作表：{spec['sheet']}")
                cells = list(wb[matching[0]].iter_rows())
                headers = [clean(c.value) for c in cells[0]] if cells else []
                for row in cells[1:]:
                    for i, cell in enumerate(row):
                        if cell.data_type == 'f':
                            raise ValueError(f"{spec['sheet']} {cell.coordinate} 不接受公式，请粘贴值")
                        if i < len(headers) and headers[i] in ('数值', 'Numeric value', 'value') and cell.value is not None and '%' in cell.number_format:
                            raise ValueError(f"{spec['sheet']} {cell.coordinate} 请使用普通数字85表示85%，不要使用Excel百分比格式")
                result[name] = _rows(name, [[c.value for c in row] for row in cells])
        finally:
            wb.close()
    else:
        raise ValueError("输入必须为模板 xlsx、含两张 CSV 的目录或标准 JSON")
    return result


def _rows(name, rows):
    if not rows:
        raise ValueError(f"{name} 缺少表头")
    spec = TABLES[name]
    labels = {f["label"]: f["key"] for f in spec["fields"]}
    labels.update({EN[f['key']][0]: f['key'] for f in spec['fields']})
    keys = {f["key"] for f in spec["fields"]}
    header = [labels.get(clean(x), clean(x)) for x in rows[0]]
    active = [x for x in header if x]
    if len(set(active)) != len(active):
        raise ValueError(f"{name} 有重复列名")
    missing = keys - set(header)
    if missing:
        raise ValueError(f"{name} 缺少列：{sorted(missing)}；请保留完整模板表头")
    unknown = set(active) - keys
    if unknown:
        raise ValueError(f"{name} 不认识的列：{sorted(unknown)}；请把补充信息放备注")
    out = []
    for row in rows[1:]:
        if any(clean(v) not in (None, "") for v in row):
            if any(clean(v) not in (None, "") and (i >= len(header) or not header[i])
                   for i, v in enumerate(row)):
                raise ValueError(f"{name} 存在无表头的数据列")
            out.append({key: clean(row[i]) if i < len(row) else None
                        for i, key in enumerate(header) if key})
    return out


def validate(data):
    errors, warnings = [], []
    if data.get("schema_version") != VERSION:
        errors.append(f"schema_version 必须为 {VERSION}")
    for name, spec in TABLES.items():
        if not isinstance(data.get(name), list):
            errors.append(f"{name} 必须是记录列表")
            continue
        seen = set()
        for i, row in enumerate(data[name], 2):
            tag = f"{spec['sheet']} 第{i}行"
            if not isinstance(row, dict):
                errors.append(f"{tag} 必须是对象")
                continue
            unknown = set(row) - {f["key"] for f in spec["fields"]}
            if unknown:
                errors.append(f"{tag} 未知字段 {sorted(unknown)}")
            for field in spec["fields"]:
                key = field["key"]
                value = canonical_choice(field, clean(row.get(key)))
                row[key] = value
                where = f"{tag} {field['label']}"
                if value in (None, ""):
                    if field["required"]:
                        errors.append(f"{where} 必填")
                    continue
                if isinstance(value, str) and value.startswith("=") and key != "qualifier":
                    errors.append(f"{where} 不接受公式，请粘贴值")
                if field["choices"] and value not in field["choices"]:
                    errors.append(f"{where} 不在允许选项内：{value}")
                kind = field["kind"]
                if kind in ("number", "integer"):
                    try:
                        number = float(value)
                        if isinstance(value, bool) or not math.isfinite(number):
                            raise ValueError
                        if kind == "integer" and not number.is_integer():
                            raise ValueError
                        row[key] = int(number) if kind == "integer" else number
                    except (ValueError, TypeError, OverflowError):
                        errors.append(f"{where} 需要有限数值")
                elif kind == "date":
                    try:
                        if len(str(value)) != 10:
                            raise ValueError
                        date.fromisoformat(str(value))
                    except ValueError:
                        errors.append(f"{where} 必须为 YYYY-MM-DD，不要用年份补造日期")
                elif not isinstance(value, str):
                    errors.append(f"{where} 需要文本")
            key = spec["fields"][0]["key"]
            ident = str(row.get(key, ""))
            if ident in seen:
                errors.append(f"{tag} 重复编号 {ident}")
            seen.add(ident)
    if errors:
        return {"errors": errors, "warnings": warnings}
    contexts = {r["context_id"]: r for r in data["contexts"]}
    for row in contexts.values():
        tag = row["context_id"]
        if row["catalysis_mode"] == "酶催化" and not row.get("scaffold"):
            errors.append(f"{tag} 酶催化需要蛋白骨架；原文未报告可填 未报告")
        if row.get("mechanism") and (not row.get("mechanism_evidence") or not row.get("mechanism_locator")):
            errors.append(f"{tag} 机理描述需要证据类型和出处")
        if row.get("time_h") is not None and row.get("time_h") != "" and row["time_h"] < 0:
            errors.append(f"{tag} 时间不能为负")
        if not row.get("first_public_date") or not row.get("date_evidence"):
            warnings.append(f"{tag} 最早公开日期/核查来源待补，不可直接用于严格时间划分")
    assays = {}
    for row in data["experiments"]:
        tag = row["measurement_id"]
        if row["context_id"] not in contexts:
            errors.append(f"{tag} 未找到体系编号 {row['context_id']}")
        val = row.get("value")
        if val not in (None, ""):
            if not row.get("qualifier") or not row.get("unit"):
                errors.append(f"{tag} 有数值时必须填写比较符和单位")
            if row["metric"] in ("yield", "conversion", "ee") and row.get("unit") != "%":
                errors.append(f"{tag} yield/conversion/ee 的单位必须为 %")
            if row.get("unit") == "%" and not (-100 if row["metric"] == "ee" else 0) <= val <= 100:
                # Relative activity can exceed 100% of its reference.
                if row["metric"] != "relative_activity":
                    errors.append(f"{tag} 百分数超出范围；85% 请填 85")
            if row["metric"] == "relative_activity" and val < 0:
                errors.append(f"{tag} 相对活性不能为负")
            if row["metric"] == "TTN" and (val < 0 or row.get("unit") != "turnovers"):
                errors.append(f"{tag} TTN 必须非负且单位为 turnovers")
            if row["metric"] == "product_detected" and (val not in (0, 1) or row.get("unit") != "boolean" or row.get("qualifier") != "="):
                errors.append(f"{tag} product_detected 用数值 0/1、单位 boolean、比较符 =")
            if row["metric"] == "product_detected" and ((val == 0 and row["outcome"] == "检出") or (val == 1 and row["outcome"] == "未检出")):
                errors.append(f"{tag} 检出状态与数值矛盾")
        elif not row.get("raw_result"):
            errors.append(f"{tag} 数值为空时保留原文结果，如 n.d. / 未报告 / 95:5")
        if row.get("replicates") not in (None, "") and row["replicates"] < 1:
            errors.append(f"{tag} 重复次数必须大于零")
        if row.get("structure_status") == "已核对":
            if not row.get("reaction_smiles"):
                errors.append(f"{tag} 已核对结构必须有反应 SMILES")
            else:
                from .chemistry import canonical_reaction
                try:
                    canonical_reaction(row["reaction_smiles"])
                except ValueError as exc:
                    errors.append(f"{tag} {exc}")
            if row.get("product_role") not in ("观察产物", "测试目标"):
                errors.append(f"{tag} 已核对结构需说明产物角色")
        else:
            warnings.append(f"{tag} 结构待补/待核对，暂不参与检索")
        if row["outcome"] == "未检出" and row.get("product_role") == "观察产物":
            errors.append(f"{tag} 未检出不能标记为观察产物")
        # Multiple measurements for one assay must describe the same chemistry and variant.
        identity = tuple(row.get(k) or "" for k in ("context_id", "compound_ids", "variant", "stage"))
        previous = assays.setdefault(row["assay_id"], identity)
        if identity != previous:
            errors.append(f"{tag} 同一实验编号的体系/化合物/变体/阶段不一致")
    return {"errors": errors, "warnings": warnings}


def write_json(path, obj):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(obj, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
