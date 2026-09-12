import csv
from copy import deepcopy
import json
from pathlib import Path
import pytest
from bioreaction_atlas.curation import read_dataset, validate, write_json
from bioreaction_atlas.localization import localized_metadata, CHOICES
from bioreaction_atlas.merge import merge_datasets
from bioreaction_atlas.schema import TABLES, VERSION


def fixture():
    return {'schema_version': VERSION, 'contexts': [{
        'context_id':'WK_C1','discovery_id':'WK_D1','source':'fictional test source','source_locator':'fixture only',
        'reaction_name':'test','catalysis_mode':'Enzymatic','scaffold':'fixture','metal_cofactor':'not reported',
        'conditions_raw':'fixture only','record_status':'Needs review'}], 'experiments':[{
        'measurement_id':'WK_M1','assay_id':'WK_A1','context_id':'WK_C1','source_locator':'fixture only',
        'compound_ids':'A -> B','stage':'Initial activity','outcome':'Not detected','metric':'product_detected',
        'raw_result':'n.d.','structure_status':'Structure pending','product_role':'Tested target'}]}


def test_english_csv_normalizes_and_merges(tmp_path):
    english = fixture()
    schema = localized_metadata('en')
    en_dir = tmp_path/'english'
    en_dir.mkdir()
    for name, spec in schema['tables'].items():
        with (en_dir/f'{name}.csv').open('w', newline='') as f:
            w = csv.writer(f)
            w.writerow([r['label'] for r in spec['fields']])
            for row in english[name]:
                w.writerow([row.get(r['key'],'') for r in spec['fields']])
    imported = read_dataset(en_dir)
    assert not validate(imported)['errors']
    assert imported['contexts'][0]['catalysis_mode'] == '酶催化'
    assert imported['experiments'][0]['outcome'] == '未检出'
    # Same logical record in Chinese is deduplicated after language normalization.
    zh_file = tmp_path/'chinese.json'
    write_json(zh_file, imported)
    merged = merge_datasets([en_dir, zh_file])
    assert len(merged['contexts']) == 1
    assert len(merged['experiments']) == 1
    assert merged['imports'][1]['identical_rows_deduplicated'] == 2


def test_conflict_never_overwrites_and_prefix_preserves_foreign_keys(tmp_path):
    a, b = fixture(), fixture()
    b['contexts'][0]['conditions_raw'] = 'different source value'
    left, right = tmp_path/'a.json', tmp_path/'b.json'
    write_json(left,a)
    write_json(right,b)
    with pytest.raises(ValueError, match='Conflicting'):
        merge_datasets([left,right])
    merged = merge_datasets([left,right], ['EM','WK'])
    assert len(merged['contexts']) == 2
    assert not validate(merged)['errors']
    assert merged['experiments'][1]['context_id'] == 'WK_WK_C1'


def test_all_fields_and_enums_have_english_translations():
    en = localized_metadata('en')
    for key, spec in TABLES.items():
        assert [r['key'] for r in spec['fields']] == [r['key'] for r in en['tables'][key]['fields']]
    assert not any('\u4e00' <= char <= '\u9fff' for char in json.dumps(en, ensure_ascii=False))


def test_english_workbook_native_features():
    from openpyxl import load_workbook
    file = Path('outputs/bioreaction_atlas_v02/Nonnatural_Enzyme_Reaction_Collection_EN.xlsx')
    assert file.exists()
    data = read_dataset(file)
    assert not data['contexts'] and not data['experiments']
    wb = load_workbook(file)
    assert wb.sheetnames == ['Reaction contexts','Measurements','Instructions','Field dictionary','Examples']
    for name, spec in localized_metadata('en')['tables'].items():
        s = wb[spec['sheet']]
        assert [c.value for c in s[1]] == [f['label'] for f in spec['fields']]
        assert s.freeze_panes == 'A2'
        assert len(s.data_validations.dataValidation) == sum(bool(f['choices']) for f in spec['fields'])
    assert all(c.data_type != 'f' for s in wb for row in s for c in row)
    wb.close()
