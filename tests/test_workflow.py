from copy import deepcopy
import pytest
from bioreaction_atlas.curation import validate, read_dataset, _rows
from bioreaction_atlas.chemistry import canonical_reaction, fingerprint, cosine
from bioreaction_atlas.retrieval import build_index, query
from bioreaction_atlas.schema import VERSION, TABLES


@pytest.fixture
def data():
    # Synthetic test fixtures only; never installed in the scientific data directory.
    return {"schema_version": VERSION, "contexts": [{
        "context_id": "TEST_C1", "discovery_id": "TEST_D1", "source": "synthetic fixture",
        "source_locator": "test only", "reaction_name": "test oxidation",
        "catalysis_mode": "非酶催化", "metal_cofactor": "test only", "conditions_raw": "test only",
        "record_status": "已核对", "first_public_date": "2020-01-02", "date_evidence": "test only",
    }], "experiments": [{
        "measurement_id": "TEST_M1", "assay_id": "TEST_A1", "context_id": "TEST_C1",
        "source_locator": "test only", "compound_ids": "TEST_1 -> TEST_2", "stage": "未说明",
        "outcome": "检出", "metric": "yield", "value": 85, "qualifier": "=", "unit": "%",
        "reaction_smiles": "CCO>>CC=O", "structure_status": "已核对", "product_role": "观察产物",
    }]}


def test_round_trip_retrieval_and_multiple_metrics(data):
    second = deepcopy(data['experiments'][0])
    second.update(measurement_id='TEST_M2', metric='ee', value=90)
    data['experiments'].append(second)
    index = build_index(data)
    assert len(index['entries']) == 1
    hit = query(index, 'OCC>>O=CC')['hits'][0]
    assert hit['reaction_similarity'] == pytest.approx(1)
    assert len(hit['measurements']) == 2
    assert hit['context']['source'] == 'synthetic fixture'


def test_zero_is_preserved_but_missing_is_not_imputed(data):
    row = data['experiments'][0]
    row.update(value=0, outcome='未检出', product_role='测试目标')
    assert not validate(data)['errors']
    assert row['value'] == 0
    row.update(value=None, raw_result='n.d.')
    assert not validate(data)['errors']
    assert row['value'] is None
    row['raw_result'] = ''
    assert validate(data)['errors']


def test_examples_never_enter_index_even_with_draft_override(data):
    data['contexts'][0]['record_status'] = '示例'
    assert not build_index(data, include_drafts=True)['entries']


def test_unverified_structures_excluded(data):
    row = data['experiments'][0]
    row.update(reaction_smiles='', structure_status='待补结构')
    assert not validate(data)['errors']
    assert not build_index(data)['entries']


@pytest.mark.parametrize('change', [
    {'context_id': 'missing'}, {'value': 120}, {'value': float('nan')},
    {'value': True}, {'qualifier': ''}, {'unit': 'ratio'},
    {'outcome': '未检出', 'product_role': '观察产物'},
    {'reaction_smiles': 'invalid>>smiles'}, {'replicates': 0},
])
def test_bad_measurements_fail(data, change):
    data['experiments'][0].update(change)
    assert validate(data)['errors']


def test_bounds_are_not_exact_values(data):
    data['experiments'][0].update(value=1, qualifier='<')
    index = build_index(data)
    assert index['entries'][0]['measurements'][0]['qualifier'] == '<'


def test_duplicate_id_and_assay_conflict(data):
    data['experiments'].append(deepcopy(data['experiments'][0]))
    assert validate(data)['errors']
    data['experiments'][1].update(measurement_id='TEST_M2', variant='different')
    assert any('不一致' in e for e in validate(data)['errors'])


def test_atom_map_order_and_stereochemistry():
    assert canonical_reaction('[CH3:1][CH2:2][OH:3]>>[CH3:1][CH:2]=[O:3]') == 'CCO>>CC=O'
    assert fingerprint('O.CCO>>CC=O.O') == fingerprint('CCO>>CC=O')
    a, b = fingerprint('CCO>>CC=O'), fingerprint('CC=O>>CCO')
    assert cosine(a, b) == pytest.approx(-1)
    assert canonical_reaction('C[C@H](O)F>>C[C@@H](O)F').split('>>')[0] != canonical_reaction('C[C@H](O)F>>C[C@@H](O)F').split('>>')[1]
    assert fingerprint('C[C@H](O)F>>C[C@@H](O)F')


def test_filters_and_version_mismatch(data):
    index = build_index(data)
    assert query(index, 'CCO>>CC=O', before='2020-01-02')['hits'] == []
    assert len(query(index, 'CCO>>CC=O', before='2020-01-03')['hits']) == 1
    assert not query(index, 'CCO>>CC=O', exclude_source='synthetic fixture')['hits']
    assert not query(index, 'CCO>>CC=O', catalysis_mode='酶催化')['hits']
    index['parameters'] = {'different': True}
    with pytest.raises(ValueError):
        query(index, 'CCO>>CC=O')


def test_unknown_or_removed_columns_fail():
    with pytest.raises(ValueError):
        _rows('contexts', [['体系编号'], ['one']])


def test_csv_import_to_cli_query(data, tmp_path):
    import csv
    import json
    from bioreaction_atlas.cli import main
    for name, spec in TABLES.items():
        with (tmp_path / f'{name}.csv').open('w', encoding='utf-8-sig', newline='') as f:
            writer = csv.writer(f)
            writer.writerow([field['label'] for field in spec['fields']])
            for row in data[name]:
                writer.writerow([row.get(field['key'], '') for field in spec['fields']])
    report, normalized = tmp_path/'report.json', tmp_path/'normalized.json'
    index, results = tmp_path/'index.json', tmp_path/'results.json'
    assert main(['validate', str(tmp_path), '--out', str(report), '--normalized', str(normalized)]) == 0
    assert main(['index', str(normalized), '--out', str(index)]) == 0
    assert main(['query', str(index), 'CCO>>CC=O', '--out', str(results)]) == 0
    assert json.loads(results.read_text())['hits'][0]['reaction_similarity'] == pytest.approx(1)


def test_delivered_blank_workbook():
    from pathlib import Path
    file = Path('outputs/bioreaction_atlas_v01/非天然酶反应收集模板.xlsx')
    if not file.exists():
        pytest.skip('Workbook not yet built')
    data = read_dataset(file)
    assert data['contexts'] == [] and data['experiments'] == []
    assert not validate(data)['errors']
    from openpyxl import load_workbook
    wb = load_workbook(file)
    for name, spec in TABLES.items():
        sheet = wb[spec['sheet']]
        assert [c.value for c in sheet[1]] == [f['label'] for f in spec['fields']]
        assert sheet.freeze_panes == 'A2'
        assert len(sheet.tables) == 1
        assert len(sheet.data_validations.dataValidation) == sum(bool(f['choices']) for f in spec['fields'])
    # The literal equals sign in the qualifier dropdown must survive Excel export.
    assert any('"=,' in (dv.formula1 or '') for dv in wb['实验结果'].data_validations.dataValidation)
    wb.close()
